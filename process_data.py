import geopandas as gpd
import openpyxl
import json
import os
import shutil

SRC = 'D:/vs code/USAP ANN PROJECT/extracted3'
EXCEL = 'D:/vs code/USAP ANN PROJECT/USAP_Fire_Risk_Index.xlsx'
OUT = 'D:/vs code/USAP ANN PROJECT/data'

if os.path.exists(OUT):
    shutil.rmtree(OUT)
os.makedirs(OUT, exist_ok=True)

# --- 1. Read Excel Risk_Index data ---
wb = openpyxl.load_workbook(EXCEL, data_only=True)
ws = wb['Risk_Index']

# Get header row
headers = [str(cell.value) for cell in ws[1] if cell.value is not None]
print(f'Excel headers ({len(headers)}): {headers}')

# Read all enterprise rows
enterprises = []
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    if row[0] is None:
        break
    enterprises.append(row)

print(f'Total enterprises in Excel: {len(enterprises)}')

# --- 2. Read survey locations ---
survey_f = os.path.join(SRC, 'Survey Locations .gpkg')
survey_gdf = gpd.read_file(survey_f)
print(f'Survey points: {len(survey_gdf)}')

# --- 3. Merge survey points with Excel data (by row index) ---
n_to_merge = min(len(survey_gdf), len(enterprises))

# Add all Excel columns as properties to survey features
for i in range(n_to_merge):
    for j, header in enumerate(headers):
        val = enterprises[i][j] if j < len(enterprises[i]) else None
        if isinstance(val, float):
            val = round(val, 4)
        survey_gdf.at[i, header] = val

# Drop fields that are empty/unnecessary
survey_gdf = survey_gdf.drop(columns=['Lat', 'Long'], errors='ignore')

# Save survey points with risk data
survey_out = os.path.join(OUT, 'survey_points.geojson')
survey_gdf.to_file(survey_out, driver='GeoJSON')
print(f'Saved survey points: {survey_out}')

# --- 4. Convert feature layers to GeoJSON ---
layers_to_export = {
    'Study Area.gpkg': 'study_area.geojson',
    'Residential.gpkg': 'residential.geojson',
    'Commercial.gpkg': 'commercial.geojson',
    'Com + Res .gpkg': 'com_res.geojson',
    'Ind + Res.gpkg': 'ind_res.geojson',
    'Ind+Com.gpkg': 'ind_com.geojson',
    'G+1.gpkg': 'g_plus_1.geojson',
    'G+2.gpkg': 'g_plus_2.geojson',
    'G+3.gpkg': 'g_plus_3.geojson',
    'G+4.gpkg': 'g_plus_4.geojson',
    'Roads .gpkg': 'roads.geojson',
    'Streets.gpkg': 'streets.geojson',
    'Drain.gpkg': 'drain.geojson',
    'Railway .gpkg': 'railway.geojson',
    'Metro Line.gpkg': 'metro_line.geojson',
    'Metro Station .gpkg': 'metro_station.geojson',
    'Health Centers.gpkg': 'health_centers.geojson',
    'School.gpkg': 'schools.geojson',
    'Religious Places .gpkg': 'religious_places.geojson',
    'Public Semi Public .gpkg': 'public_semi_public.geojson',
    'Open Spaces .gpkg': 'open_spaces.geojson',
    'Ward  Boundary .gpkg': 'ward_boundary.geojson',
    'Present.gpkg': 'present.geojson',
    '!962- 80.gpkg': '1962_1980.geojson',
    '1980-2000.gpkg': '1980_2000.geojson',
    '2000-2020.gpkg': '2000_2020.geojson',
}

for src_name, out_name in layers_to_export.items():
    src_path = os.path.join(SRC, src_name)
    if not os.path.exists(src_path):
        print(f'WARNING: {src_name} not found, skipping')
        continue
    try:
        gdf = gpd.read_file(src_path)
        if len(gdf) == 0:
            print(f'Skipping empty layer: {src_name}')
            continue
        out_path = os.path.join(OUT, out_name)
        gdf.to_file(out_path, driver='GeoJSON')
        print(f'Saved {out_name} ({len(gdf)} features)')
    except Exception as e:
        print(f'ERROR skipping {src_name}: {e}')

# --- 5. Create a layer manifest JSON ---
layer_manifest = {
    "basemaps": [
        {"id": "osm", "name": "OpenStreetMap", "type": "raster"},
        {"id": "satellite", "name": "Satellite", "type": "raster"}
    ],
    "overlays": [
        {"id": "survey_points", "file": "survey_points.geojson", "name": "Survey Points (Risk)", "type": "point"},
        {"id": "study_area", "file": "study_area.geojson", "name": "Study Area", "type": "polygon"},
        {"id": "residential", "file": "residential.geojson", "name": "Residential", "type": "polygon"},
        {"id": "commercial", "file": "commercial.geojson", "name": "Commercial", "type": "polygon"},
        {"id": "com_res", "file": "com_res.geojson", "name": "Commercial + Residential", "type": "polygon"},
        {"id": "ind_res", "file": "ind_res.geojson", "name": "Industrial + Residential", "type": "polygon"},
        {"id": "ind_com", "file": "ind_com.geojson", "name": "Industrial + Commercial", "type": "polygon"},
        {"id": "g_plus_1", "file": "g_plus_1.geojson", "name": "G+1 Buildings", "type": "polygon"},
        {"id": "g_plus_2", "file": "g_plus_2.geojson", "name": "G+2 Buildings", "type": "polygon"},
        {"id": "g_plus_3", "file": "g_plus_3.geojson", "name": "G+3 Buildings", "type": "polygon"},
        {"id": "g_plus_4", "file": "g_plus_4.geojson", "name": "G+4 Buildings", "type": "polygon"},
        {"id": "roads", "file": "roads.geojson", "name": "Major Roads", "type": "line"},
        {"id": "streets", "file": "streets.geojson", "name": "Streets", "type": "line"},
        {"id": "drain", "file": "drain.geojson", "name": "Drain", "type": "line"},
        {"id": "railway", "file": "railway.geojson", "name": "Railway", "type": "line"},
        {"id": "metro_line", "file": "metro_line.geojson", "name": "Metro Line", "type": "line"},
        {"id": "metro_station", "file": "metro_station.geojson", "name": "Metro Station", "type": "point"},
        {"id": "health_centers", "file": "health_centers.geojson", "name": "Health Centers", "type": "point"},
        {"id": "schools", "file": "schools.geojson", "name": "Schools", "type": "point"},
        {"id": "religious_places", "file": "religious_places.geojson", "name": "Religious Places", "type": "point"},
        {"id": "public_semi_public", "file": "public_semi_public.geojson", "name": "Public/Semi-Public", "type": "polygon"},
        {"id": "ward_boundary", "file": "ward_boundary.geojson", "name": "Ward Boundary", "type": "polygon"}
    ],
    "risk_colors": {
        "Severe": "#D32F2F",
        "High": "#FF9800",
        "Moderate": "#FFC107",
        "Low": "#4CAF50"
    }
}

with open(os.path.join(OUT, 'manifest.json'), 'w') as f:
    json.dump(layer_manifest, f, indent=2)

print('\nDone! All data processed.')
print(f'Output directory: {OUT}')
